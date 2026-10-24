from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy
from pathlib import Path

src=Document(r'C:\Users\Parth\Downloads\Literature_Review_Paper_IEEE_5pages (1).docx')
out=Path(r'C:\Users\Parth\OneDrive\Desktop\ManagementSystem\Literature_Review_Paper_IEEE_Expanded.docx')
doc=Document(out)

def find(d, starts):
    return next(p for p in d.paragraphs if p.text.startswith(starts))
def rpr(run):
    return deepcopy(run._r.rPr) if run is not None and run._r.rPr is not None else None
def set_split(p, lead, body, lead_rpr, body_rpr, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
    for run in list(p.runs): p._p.remove(run._r)
    a=p.add_run(lead)
    if lead_rpr is not None: a._r.insert(0,deepcopy(lead_rpr))
    b=p.add_run(body)
    if body_rpr is not None: b._r.insert(0,deepcopy(body_rpr))
    p.alignment=alignment

def insert_before(anchor,text,template):
    el=OxmlElement('w:p')
    if template._p.pPr is not None: el.append(deepcopy(template._p.pPr))
    anchor._p.addprevious(el)
    p=Paragraph(el,anchor._parent)
    p.style=template.style
    run=p.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None: run._r.insert(0,deepcopy(template.runs[0]._r.rPr))
    return p

# Restore inline lead-in styling with regular, readable body text.
for source_prefix, final_prefix, label in [
    ('1) Domain focus:', '1) Operating context:', '1) Operating context:'),
    ('2) Hardware dependence:', '2) Technology choice:', '2) Technology choice:'),
    ('3) Lack of integration:', '3) Workflow focus:', '3) Workflow focus:'),
    ('4) Limited evaluation:', '4) Evaluation:', '4) Evaluation:')]:
    original=find(src,source_prefix)
    current=find(doc,final_prefix)
    normal_run=original.runs[1] if len(original.runs)>1 else original.runs[0]
    lead_run=original.runs[0] if original.runs else None
    body=current.text[len(label):].lstrip()
    set_split(current,label+' ',body,rpr(lead_run),rpr(normal_run))

orig_conclusion=find(src,'Conclusion.')
conclusion=find(doc,'Conclusion. This review')
set_split(conclusion,'Conclusion. ',conclusion.text[len('Conclusion. '):],rpr(orig_conclusion.runs[0]),rpr(orig_conclusion.runs[1]))
orig_future=find(src,'Future Scope.')
future=find(doc,'Future Scope. Future evaluation')
set_split(future,'Future Scope. ',future.text[len('Future Scope. '):],rpr(orig_future.runs[0]),rpr(orig_future.runs[1]))

# Normalize the deployment recommendation paragraph as body copy, not an italic lead-in.
security=find(doc,'Before deployment,')
normal=find(src,'The work followed six steps:')
security.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for run in list(security.runs): security._p.remove(run._r)
r=security.add_run(security.text if False else 'Before deployment, the application should be assessed for password hashing, server-side authorization on every protected operation, parameterized database access, input validation, session and cookie settings, error handling, backup and restore, and transport security. These are review recommendations only; the supplied test record does not demonstrate that they are implemented or security-tested.')
if normal.runs and normal.runs[0]._r.rPr is not None: r._r.insert(0,deepcopy(normal.runs[0]._r.rPr))

# Keep the split-prone workflow-focus item together and add the missing Section III heading.
find(doc,'3) Workflow focus:').paragraph_format.keep_together=True
comparison_summary=find(doc,'Table I summarizes eight representative papers')
comparison_heading=find(doc,'IV. PROJECT POSITION')
insert_before(comparison_summary,'III. COMPARISON OF THE STUDIES',comparison_heading)

doc.save(out)
print(out)
