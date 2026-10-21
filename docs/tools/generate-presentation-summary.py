from pathlib import Path

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_styled_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Style Header
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], "1E293B") # Dark slate
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.name = "Calibri"
            run.font.size = Pt(10)

    # Style Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=150, right=150)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(30, 41, 59)

    # Apply column widths if provided
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    # Set table borders
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        '<w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        '<w:insideV w:val="none"/>'
        '<w:left w:val="none"/>'
        '<w:right w:val="none"/>'
        '</w:tblBorders>'
    )
    tblPr.append(tblBorders)
    
    doc.add_paragraph() # Spacing
    return table

def add_callout(doc, text, title="KEY HIGHLIGHT", border_color="0284C7", bg_color="F0F9FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        '<w:top w:val="none"/>'
        '<w:bottom w:val="none"/>'
        '<w:right w:val="none"/>'
        '</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = RGBColor(2, 132, 199)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph()

def build_document():
    doc = docx.Document()
    
    # Page Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)

    # Base Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run("DECOSTONE MATERIAL MANAGEMENT SYSTEM (MMS)")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.name = 'Calibri'
    title_run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    sub_run = sub_p.add_run("Comprehensive Technical Documentation, System Explanation & Presentation Summary")
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(2, 132, 199) # Sky 600

    add_callout(
        doc,
        "Enterprise-grade multi-warehouse material management and inventory movement system for architectural tiles, marble, granite, kitchen sinks, and sanitary ware. Enforces ACID transactions, negative stock prevention, and multi-tier RBAC.",
        "PROJECT OVERVIEW & PRESENTATION SUMMARY"
    )

    def add_heading_1(text):
        h = doc.add_heading(level=1)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(15)
        r.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
        return h

    def add_heading_2(text):
        h = doc.add_heading(level=2)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(12)
        r.bold = True
        r.font.color.rgb = RGBColor(51, 65, 85)
        return h

    # Section 1
    add_heading_1("1. Executive Summary")
    p = doc.add_paragraph("DecoStone Material Management System (MMS) is a centralized, real-time web inventory management platform designed to track architectural and sanitary materials across four distinct warehouses (Godown 1, Godown 2, Godown 3, and Godown 4).")
    p = doc.add_paragraph("The platform eliminates traditional manual warehouse errors, enforces database-level ACID transaction integrity, prevents negative stock situations, automates minimum threshold safety alerts, and provides comprehensive audit trails alongside 1-click PDF and Excel reporting.")

    # Section 2
    add_heading_1("2. Problem Statement & Business Context")
    add_heading_2("Traditional Warehouse Challenges:")
    b1 = doc.add_paragraph(style='List Bullet')
    b1.add_run("Dispersed Inventory: ").bold = True
    b1.add_run("Materials stored across multiple warehouses lack real-time cross-visibility, resulting in double-selling or stockouts.")
    
    b2 = doc.add_paragraph(style='List Bullet')
    b2.add_run("Manual Ledger Discrepancies: ").bold = True
    b2.add_run("Spreadsheet-based and paper-based tracking allow negative balances, unrecorded breakages, and inaccurate counts.")
    
    b3 = doc.add_paragraph(style='List Bullet')
    b3.add_run("Untracked Inter-Godown Transfers: ").bold = True
    b3.add_run("Moving materials between warehouses frequently creates phantom stock or duplicated figures.")
    
    b4 = doc.add_paragraph(style='List Bullet')
    b4.add_run("Uncontrolled Access: ").bold = True
    b4.add_run("Lack of role segmentation leads to unauthorized dispatches or accidental catalog modifications.")

    # Section 3
    add_heading_1("3. Core Project Objectives")
    headers_obj = ["Objective", "System Implementation", "Business Impact"]
    rows_obj = [
        ["Independent Multi-Godown Tracking", "Real-time stock ledger per Godown (1-4) per SKU", "Accurate warehouse-level stock visibility"],
        ["Atomic Material Transfers", "Single database transaction with TRF-YYYY-NNNN code", "Zero duplicate or lost stock during relocation"],
        ["Negative Stock Prevention", "Database CHECK constraint & API validation", "Guaranteed physical stock compliance"],
        ["Role-Based Access Control", "Granular Admin, Manager, Supervisor tiers", "Secure data boundaries & operational safety"],
        ["Automated Low Stock Center", "Live threshold calculation (Quantity <= Safety Limit)", "Eliminates project delays and sudden shortages"],
        ["Multi-Format Reports", "1-Click PDF (jsPDF) & Excel (.xlsx) exports", "Instant stakeholder audits & operational reviews"]
    ]
    create_styled_table(doc, headers_obj, rows_obj, [1.8, 2.5, 2.2])

    # Section 4
    add_heading_1("4. Technical Architecture & Stack")
    p = doc.add_paragraph("The system follows a modern, decoupled client-server architecture:")
    headers_tech = ["Layer", "Technology", "Role & Functionality"]
    rows_tech = [
        ["Frontend UI", "React 19 + Vite", "Fast, reactive SPA with component-driven architecture"],
        ["Styling & Design", "Vanilla CSS + Glassmorphism", "Slate/stone architectural theme, responsive layouts"],
        ["Data Visuals", "Recharts", "Interactive pie charts, inflow/outflow velocity area charts"],
        ["Document Engine", "jsPDF & SheetJS (XLSX)", "Client-side export of formatted reports and tables"],
        ["Backend REST API", "Node.js + Express.js", "Stateless REST endpoints with modular services and middleware"],
        ["Database", "SQLite (better-sqlite3)", "Embedded ACID database with Write-Ahead Logging (WAL mode)"],
        ["Security & Auth", "JWT + bcryptjs", "Stateless JSON Web Tokens with encrypted password hashing"]
    ]
    create_styled_table(doc, headers_tech, rows_tech, [1.5, 2.2, 2.8])

    # Section 5
    add_heading_1("5. Key Functional Modules")
    add_heading_2("5.1 Executive & Warehouse Dashboards")
    doc.add_paragraph("Dynamic KPI cards (Total Stock Value, Active SKUs, Low Stock Count, Today's Movement) and dedicated Godown 1–4 views showing live inventory status (🟢 In Stock, 🟠 Low Stock, 🔴 Out of Stock).")

    add_heading_2("5.2 Stock In & Stock Out Operations")
    doc.add_paragraph("Stock IN increments warehouse balance with Supplier and PO reference tracking. Stock OUT securely decrements stock with customer and invoice references, guarded against negative balances.")

    add_heading_2("5.3 Atomic Inter-Godown Transfers")
    doc.add_paragraph("Transfers move inventory between warehouses using atomic transactions. A unique TRF-YYYY-NNNN tracking number is generated. If source deduction fails, destination increment is automatically rolled back, keeping total company stock invariant.")

    add_heading_2("5.4 Stock Reconciliation & Adjustments")
    doc.add_paragraph("Enables managers and admins to record physical inventory counts against system counts, computing surplus/deficits and capturing mandatory audit reasons (e.g. transit breakage, counting adjustments).")

    add_heading_2("5.5 Low Stock Warning Center & Audit Trail")
    doc.add_paragraph("Monitors all materials against safety minimum thresholds and logs every login, catalog change, and stock movement in an immutable audit ledger.")

    # Section 6
    add_heading_1("6. Stock Calculation & Data Integrity Logic")
    add_callout(
        doc,
        "Current Stock = Opening Stock + Stock IN - Stock OUT + Transfers IN - Transfers OUT ± Adjustments\n\nGuaranteed Properties:\n1. Non-negative stock invariant (Quantity >= 0)\n2. Total Company Stock is conserved during inter-warehouse transfers\n3. Atomicity: Multi-warehouse updates succeed together or fail together.",
        "MATHEMATICAL STOCK INVARIANT FORMULA",
        border_color="10B981",
        bg_color="F0FDF4"
    )

    # Section 7
    add_heading_1("7. Role-Based Access Control (RBAC) Matrix")
    headers_rbac = ["Module / Feature", "Admin", "Manager", "Supervisor"]
    rows_rbac = [
        ["User & Role Management", "✅ Full Control", "❌ No Access", "❌ No Access"],
        ["Godown Settings & Renaming", "✅ Full Control", "❌ No Access", "❌ No Access"],
        ["Category Management", "✅ Full Control", "👁️ View Only", "❌ No Access"],
        ["Material Catalog Management", "✅ Add/Edit/Delete", "✅ Add/Edit", "❌ No Access"],
        ["Multi-Godown Overview", "✅ All Godowns", "✅ All Godowns", "🔒 Assigned Godown"],
        ["Stock IN / Stock OUT", "✅ Full Access", "✅ Full Access", "🔒 If Assigned & Permitted"],
        ["Inter-Godown Transfers", "✅ Execute & View", "✅ Execute & View", "👁️ View Status Only"],
        ["Stock Adjustments", "✅ Full Access", "🔒 Authorized Only", "❌ No Access"],
        ["Low Stock Center & Alerts", "✅ Full Access", "✅ Full Access", "❌ No Access"],
        ["PDF & Excel Export Reports", "✅ Full Access", "✅ Full Access", "❌ No Access"],
        ["System Security Audit Logs", "✅ Full Access", "❌ No Access", "❌ No Access"]
    ]
    create_styled_table(doc, headers_rbac, rows_rbac, [2.5, 1.3, 1.3, 1.4])

    # Section 8
    add_heading_1("8. Slide-by-Slide Presentation Guide")
    headers_slides = ["Slide #", "Slide Title", "Key Content & Talking Points"]
    rows_slides = [
        ["Slide 1", "Title & Introduction", "Project Name: DecoStone MMS. Problem overview & core purpose."],
        ["Slide 2", "Industry Problem Statement", "Multi-warehouse challenges in tiles/marble, manual error risks."],
        ["Slide 3", "Proposed Solution", "Centralized web platform, real-time tracking, multi-godown sync."],
        ["Slide 4", "System Architecture & Stack", "React 19, Express.js REST API, SQLite WAL Mode, JWT auth."],
        ["Slide 5", "Core Functional Modules", "Godown 1-4, SKU Catalog, Stock In/Out, Transfers, Adjustments."],
        ["Slide 6", "Atomic Transfers & Integrity", "TRF-YYYY-NNNN tracking, ACID rollback, negative stock guard."],
        ["Slide 7", "Role-Based Access Control", "Admin, Manager, Supervisor tiers with backend security."],
        ["Slide 8", "Analytics & Reporting Engine", "Dynamic Recharts, 1-click PDF & Excel (.xlsx) exports."],
        ["Slide 9", "Live System Demonstration", "Step-by-step workflow: Stock IN -> Transfer -> Guard Check -> Export."],
        ["Slide 10", "Scalability & Future Scope", "Barcode/QR scanning, automated POs, supplier B2B portals."],
        ["Slide 11", "Conclusion & Q/A", "Project impact summary and opening floor for questions."]
    ]
    create_styled_table(doc, headers_slides, rows_slides, [1.0, 2.0, 3.5])

    # Section 9
    add_heading_1("9. Live Demonstration & Verification Script")
    steps = [
        ("1. Admin Login & Overview", "Log in as Admin (admin / admin123). Showcase real-time KPI cards, Godown distribution charts, and inventory summary."),
        ("2. Catalog Exploration", "Open Categories & Materials. Demonstrate SKU attributes (dimensions, box units, safety minimum levels)."),
        ("3. Stock IN Operation", "Record receipt of 50 boxes into Godown 1 with PO number. Verify immediate ledger and balance update."),
        ("4. Inter-Warehouse Transfer", "Transfer 20 boxes from Godown 1 to Godown 2. Verify generation of TRF tracking ID and invariant total company stock."),
        ("5. Negative Stock Guard Test", "Attempt to dispatch 9,999 boxes from Godown 3. Highlight system rejection: 'Insufficient stock available.'"),
        ("6. Low Stock & Adjustment", "Inspect Low Stock Center. Reconcile inventory difference with audit justification."),
        ("7. Report Generation", "Export multi-godown ledger to PDF and Excel spreadsheets with formatted tables."),
        ("8. Supervisor Isolation", "Log in as Supervisor 1 to show restricted single-warehouse view.")
    ]
    for stitle, sdesc in steps:
        bp = doc.add_paragraph(style='List Bullet')
        bp.add_run(f"{stitle}: ").bold = True
        bp.add_run(sdesc)

    # Section 10
    add_heading_1("10. Future Enhancements")
    doc.add_paragraph("• Mobile Barcode & QR Code Scanning: Direct pallet scanning via mobile device cameras.\n• Supplier & Customer B2B Portal: Automated PO acknowledgment and client dispatch tracking.\n• Automated Reorder Triggers: Direct notification and purchase order drafting when stock drops below safety levels.\n• Multi-Branch Expansion: Regional scaling across multiple geographic locations and central distribution hubs.")

    # Footer notice
    doc.add_paragraph()
    ft_p = doc.add_paragraph()
    ft_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ft_run = ft_p.add_run("— End of Documentation — Prepared for DecoStone MMS Project Presentation & Evaluation —")
    ft_run.font.italic = True
    ft_run.font.size = Pt(8.5)
    ft_run.font.color.rgb = RGBColor(148, 163, 184)

    output_path = Path(__file__).resolve().parents[1] / "presentation-summary.docx"
    doc.save(output_path)
    print("Document successfully created at:", output_path)

if __name__ == "__main__":
    build_document()
